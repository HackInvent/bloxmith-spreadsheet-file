"""XLSX support through openpyxl, imported lazily so CSV keeps working without it.

The package is inspected before openpyxl opens it: bounded ZIP entries and
expansion, no XML document type declarations, explicit macro detection. Workbooks
are opened read-only with cached formula values; external links are never loaded.
"""

from __future__ import annotations

import datetime as dt
import importlib.util
import io
import math
import re
import zipfile
import zlib

from .logic import (SAFE_INTEGER, XLSX_CELL_CHARS, XLSX_ILLEGAL, XLSX_MAX_ROWS, SpreadsheetError, fail, tabulate)

ARCHIVE_ERRORS = (zipfile.BadZipFile, zipfile.LargeZipFile, zlib.error, NotImplementedError, OSError, EOFError,
                  RuntimeError, ValueError)

MAX_ENTRIES = 10000
MAX_TOTAL_EXPANDED = 1024 * 1024 * 1024
MAX_PART_EXPANDED = 128 * 1024 * 1024
MAX_RATIO = 1000
PROLOG_WINDOW = 65536
EXTERNAL_LINK = re.compile(r"xl/externallinks/externallink\d+\.xml")
PROLOG_TOKEN = re.compile(r"\s+|<\?.*?\?>|<!--.*?-->", re.S)
DEPENDENCY = "XLSX support needs openpyxl (>=3.1,<4) in the Studio Python environment; CSV works without it."


def openpyxl_available():
    """Find the optional dependency without importing it (safe in prepare_runtime)."""
    try:
        return importlib.util.find_spec("openpyxl") is not None
    except (ImportError, ValueError):
        return False


def load_openpyxl():
    """Import openpyxl on demand and refuse versions older than the tested API."""
    try:
        import openpyxl  # optional dependency, deliberately not imported at module level
    except ImportError:
        fail("missing_dependency", DEPENDENCY)
    try:
        major, minor = (int(part) for part in str(openpyxl.__version__).split(".")[:2])
    except (AttributeError, ValueError):
        major, minor = 0, 0
    if (major, minor) < (3, 1) or major >= 4:
        fail("missing_dependency", DEPENDENCY)
    return openpyxl


def declares_dtd(head):
    """Report whether an XML part declares a DTD (or other markup declaration) before its root."""
    if head.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = head[:len(head) - len(head) % 2].decode("utf-16", errors="replace")
    else:
        text = head.decode("utf-8-sig", errors="replace")
    position = 0
    while True:
        match = PROLOG_TOKEN.match(text, position)
        if not match or match.end() == position:
            break
        position = match.end()
    rest = text[position:position + 2]
    if rest.startswith("<!"):
        return True
    if len(rest) == 2 and rest[0] == "<" and (rest[1].isalpha() or rest[1] in "_:"):
        return False
    if not text[position:].strip():
        return False
    # An unterminated declaration, comment or instruction could hide a DTD beyond the window.
    fail("malformed_file", "An XLSX XML part has an invalid or unusually long prolog.")


def inspect_package(data, allow_macros, check):
    """Bound the ZIP structure before any XML is parsed; return safe package facts."""
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        fail("unsupported_format", "Password-protected workbooks and legacy .xls files are not supported.")
    if not data.startswith(b"PK"):
        fail("malformed_file", "The file is not an XLSX (ZIP) package.")
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        infos = archive.infolist()
    except ARCHIVE_ERRORS:
        fail("malformed_file", "The XLSX package is not a readable ZIP archive.")
    try:
        names = [info.filename for info in infos]
        if len(infos) > MAX_ENTRIES or len(set(names)) != len(names):
            fail("malformed_file", "The XLSX package has too many or duplicate entries.")
        total = 0
        for info in infos:
            check()
            if info.flag_bits & 0x1:
                fail("unsupported_format", "Encrypted ZIP entries are not supported.")
            total += info.file_size
            worksheet = info.filename.lower().startswith("xl/worksheets/")
            if (not worksheet and info.file_size > MAX_PART_EXPANDED) or total > MAX_TOTAL_EXPANDED or (
                    info.file_size > 1048576 and info.file_size > MAX_RATIO * max(info.compress_size, 1)):
                fail("too_large", "The XLSX package expands beyond its safety limits.")
            if info.filename.lower().endswith((".xml", ".rels")):
                with archive.open(info) as part:
                    if declares_dtd(part.read(PROLOG_WINDOW)):
                        fail("malformed_file", "XML document type declarations are not allowed in XLSX parts.")
        lowered = [name.lower() for name in names]
        macro = "xl/vbaproject.bin" in lowered
        if macro and not allow_macros:
            fail("macro_workbook", "The workbook contains a macro project; enable allow_macro_workbooks to read its data.")
        return {"macro_workbook": macro, "external_links": sum(bool(EXTERNAL_LINK.fullmatch(name)) for name in lowered)}
    except SpreadsheetError:
        raise
    except ARCHIVE_ERRORS:
        fail("malformed_file", "The XLSX package is not a readable ZIP archive.")
    finally:
        archive.close()


def iso_datetime(value):
    """Midnight becomes a calendar date; other moments keep seconds (or milliseconds)."""
    if value.time() == dt.time(0) and value.tzinfo is None:
        return value.date().isoformat()
    return value.isoformat(timespec="milliseconds" if value.microsecond else "seconds")


def iso_duration(value):
    """Render an Excel duration as an ISO 8601 duration such as P1DT2H3M."""
    total = round(value.total_seconds() * 1000)
    sign, total = ("-" if total < 0 else ""), abs(total)
    days, rest = divmod(total, 86400000)
    hours, rest = divmod(rest, 3600000)
    minutes, rest = divmod(rest, 60000)
    seconds, millis = divmod(rest, 1000)
    clock = (f"{hours}H" if hours else "") + (f"{minutes}M" if minutes else "")
    if seconds or millis:
        clock += (f"{seconds}.{millis:03d}".rstrip("0").rstrip(".") if millis else str(seconds)) + "S"
    text = "P" + (f"{days}D" if days else "") + ("T" + clock if clock else "")
    return sign + (text if text != "P" else "PT0S")


def cell_value(value):
    """Convert openpyxl values to JSON: dates as ISO strings, error cells as their text."""
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            fail("malformed_file", "A cell holds a non-finite number.")
        return value
    if isinstance(value, dt.datetime):
        return iso_datetime(value)
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, dt.timedelta):
        return iso_duration(value)
    return str(value)


def read_xlsx(data, config, check, macro_suffix):
    """Read one worksheet as cached values; openpyxl never evaluates formulas or follows links."""
    facts = inspect_package(data, config["allow_macro_workbooks"], check)
    openpyxl = load_openpyxl()
    check()
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False, keep_vba=False)
    except Exception:  # noqa: BLE001 - openpyxl raises many types for malformed packages.
        fail("malformed_file", "The workbook structure cannot be read.")
    try:
        sheets = list(book.worksheets)
        titles = [sheet.title for sheet in sheets]
        if config["sheet"]:
            if config["sheet"] not in titles:
                listed = ", ".join(title[:31] for title in titles[:10]) + (", ..." if len(titles) > 10 else "")
                fail("invalid_input", f"Worksheet not found; available worksheets: {listed or 'none'}.")
            index = titles.index(config["sheet"]) + 1
        else:
            index = config["sheet_index"]
            if index > len(sheets):
                fail("invalid_input", f"sheet_index {index} exceeds the {len(sheets)} worksheet(s) of the workbook.")
        sheet = sheets[index - 1]
        # Never trust the declared dimension: a forged A1:XFD1048576 would pad every row.
        sheet.reset_dimensions()

        def records():
            iterator = sheet.iter_rows(values_only=True)
            while True:
                try:
                    row = next(iterator)
                except StopIteration:
                    return
                except Exception:  # noqa: BLE001 - malformed sheet XML surfaces as many exception types.
                    fail("malformed_file", "The worksheet data cannot be parsed.")
                yield [cell_value(value) for value in row]
        rows, shape = tabulate(records(), config, check, XLSX_MAX_ROWS)
        return rows, {**shape, "sheet": sheet.title, "sheet_index": index, "sheet_state": sheet.sheet_state,
                      "sheets": titles[:256], "sheets_truncated": len(titles) > 256,
                      "macro_workbook": facts["macro_workbook"] or macro_suffix,
                      "external_links": facts["external_links"], "formulas": "cached_values"}
    finally:
        book.close()


def write_xlsx(header, matrix, config, check, scalar):
    """Write one worksheet in openpyxl write-only mode with every text cell typed as literal text."""
    openpyxl = load_openpyxl()
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    counters = {"json_cells": 0, "numbers_as_text": 0}
    book = openpyxl.Workbook(write_only=True)
    book.properties.creator = "BloxSmith Spreadsheet File"
    sheet = book.create_sheet(config["sheet"] or "Sheet1")

    def text_cell(value, position, bold=False):
        if len(value) > XLSX_CELL_CHARS:
            fail("invalid_value", f"{position}: text exceeds the XLSX limit of 32,767 characters.")
        if XLSX_ILLEGAL.search(value):
            fail("invalid_value", f"{position}: text contains control characters that XLSX cannot store.")
        cell = WriteOnlyCell(sheet, value=value)
        cell.data_type = "s"  # literal text: never a formula or an error code
        if bold:
            cell.font = Font(bold=True)
        return cell

    def converted(value, position):
        value = scalar(value, counters)
        if isinstance(value, str):
            return text_cell(value, position)
        if isinstance(value, int) and not isinstance(value, bool) and abs(value) > SAFE_INTEGER:
            counters["numbers_as_text"] += 1
            return text_cell(str(value), position)
        return value

    if header is not None:
        sheet.freeze_panes = "A2"
        sheet.append([text_cell(str(name), f"Header column {index}", bold=True) for index, name in enumerate(header, 1)])
    for number, row in enumerate(matrix, 1):
        check()
        sheet.append([converted(value, f"Row {number}, column {index}") for index, value in enumerate(row, 1)])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue(), {"sheet": config["sheet"] or "Sheet1", **counters}
