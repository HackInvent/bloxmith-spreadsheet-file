"""Synthetic spreadsheet fixtures for the block suites; no personal or production file is used."""

import datetime as dt
import io
import struct
import zipfile

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
DOC_RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SHEET_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"


def minimal_xlsx(sheets, *, shared=None, extra=None, prolog='<?xml version="1.0" encoding="UTF-8"?>'):
    """Build a small hand-written package: {sheet name: sheetData XML}, optional shared strings/parts."""
    overrides = ['<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-'
                 'officedocument.spreadsheetml.sheet.main+xml"/>']
    entries, workbook_sheets, relations = {}, [], []
    for index, (name, rows) in enumerate(sheets.items(), 1):
        overrides.append(f'<Override PartName="/xl/worksheets/sheet{index}.xml" ContentType="{SHEET_TYPE}"/>')
        workbook_sheets.append(f'<sheet name="{name}" sheetId="{index}" r:id="rId{index}"/>')
        relations.append(f'<Relationship Id="rId{index}" Type="{DOC_RELS}/worksheet" Target="worksheets/sheet{index}.xml"/>')
        entries[f"xl/worksheets/sheet{index}.xml"] = f'{prolog}<worksheet xmlns="{MAIN}"><sheetData>{rows}</sheetData></worksheet>'
    if shared is not None:
        overrides.append('<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-'
                         'officedocument.spreadsheetml.sharedStrings+xml"/>')
        relations.append(f'<Relationship Id="rIdS" Type="{DOC_RELS}/sharedStrings" Target="sharedStrings.xml"/>')
        items = "".join(f"<si><t>{text}</t></si>" for text in shared)
        entries["xl/sharedStrings.xml"] = f'{prolog}<sst xmlns="{MAIN}" count="{len(shared)}" uniqueCount="{len(shared)}">{items}</sst>'
    entries["[Content_Types].xml"] = (f'{prolog}<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                                      '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                                      '<Default Extension="xml" ContentType="application/xml"/>' + "".join(overrides) + "</Types>")
    entries["_rels/.rels"] = (f'{prolog}<Relationships xmlns="{RELS}"><Relationship Id="rId1" '
                              f'Type="{DOC_RELS}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
    entries["xl/workbook.xml"] = (f'{prolog}<workbook xmlns="{MAIN}" xmlns:r="{DOC_RELS}"><sheets>'
                                  + "".join(workbook_sheets) + "</sheets></workbook>")
    entries["xl/_rels/workbook.xml.rels"] = f'{prolog}<Relationships xmlns="{RELS}">' + "".join(relations) + "</Relationships>"
    entries.update(extra or {})
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def inline(ref, text):
    """One inline-string cell."""
    return f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>'


def forge_entry(data, name, *, size=None, flags=None):
    """Patch one central-directory record (declared size or flag bits) without touching its data."""
    raw = bytearray(data)
    position = raw.find(b"PK\x01\x02")
    target = name.encode("utf-8")
    while position >= 0:
        length = struct.unpack_from("<H", raw, position + 28)[0]
        if bytes(raw[position + 46:position + 46 + length]) == target:
            if flags is not None:
                struct.pack_into("<H", raw, position + 8, flags)
            if size is not None:
                struct.pack_into("<I", raw, position + 24, size)
            return bytes(raw)
        position = raw.find(b"PK\x01\x02", position + 4)
    raise AssertionError("entry not found")


def rich_workbook(path):
    """Write a workbook with preamble rows, typed cells, a hidden sheet and a chart sheet via openpyxl."""
    import openpyxl
    book = openpyxl.Workbook()
    notes = book.active
    notes.title = "Notes"
    notes["A1"] = "This sheet is documentation"
    data = book.create_sheet("Data")
    data.append(["Quarterly export"])
    data.append([])
    data.append(["id", "day", "moment", "clock", "span", "ok", "ratio", None, "id", 2024])
    data.append([1, dt.date(2024, 1, 31), dt.datetime(2024, 1, 31, 8, 30, 5), dt.time(8, 30),
                 dt.timedelta(hours=26, minutes=3), True, 0.25, None, "first", "x"])
    data.append([])
    data.append([2, None, dt.datetime(2024, 2, 1, 0, 0), None, None, False, -1.5, None, "second", None])
    data["E4"].number_format = "[h]:mm:ss"
    hidden = book.create_sheet("Archive")
    hidden.sheet_state = "hidden"
    hidden.append(["archived"])
    hidden.append(["old row"])
    from openpyxl.chart import BarChart, Reference
    chart = BarChart()
    chart.add_data(Reference(data, min_col=7, min_row=3, max_row=6), titles_from_data=True)
    book.create_chartsheet("Chart").add_chart(chart)  # real chart sheets always carry a drawing
    book.save(path)
